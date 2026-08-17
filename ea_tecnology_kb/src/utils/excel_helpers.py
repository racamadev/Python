"""Excel export formatting helpers built on openpyxl.

Provides reusable worksheet post-processing (autofilter, frozen header
row, auto-sized columns, header styling and date formatting) so that
:class:`~src.services.export_service.ExportService` can produce a
professional-looking, multi-sheet workbook without duplicating styling
logic per sheet.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_DATE_COLUMNS_HINT = ("date",)
_MAX_COLUMN_WIDTH = 60
_MIN_COLUMN_WIDTH = 10


def style_header_row(worksheet: Worksheet) -> None:
    """Apply bold white-on-blue styling to the first (header) row.

    Args:
        worksheet: The worksheet whose header row will be styled.
    """
    for cell in worksheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def freeze_header(worksheet: Worksheet) -> None:
    """Freeze the header row so it stays visible while scrolling.

    Args:
        worksheet: The worksheet to freeze.
    """
    worksheet.freeze_panes = "A2"


def apply_autofilter(worksheet: Worksheet) -> None:
    """Enable Excel autofilter across the worksheet's used range.

    Args:
        worksheet: The worksheet to apply the autofilter to.
    """
    if worksheet.max_row < 1 or worksheet.max_column < 1:
        return
    last_col = get_column_letter(worksheet.max_column)
    worksheet.auto_filter.ref = f"A1:{last_col}{worksheet.max_row}"


def autosize_columns(worksheet: Worksheet) -> None:
    """Resize each column to fit its widest cell content.

    Args:
        worksheet: The worksheet whose columns will be resized.
    """
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        width = min(_MAX_COLUMN_WIDTH, max(_MIN_COLUMN_WIDTH, max_length + 2))
        worksheet.column_dimensions[column_letter].width = width


def apply_date_format(worksheet: Worksheet, header_names: list[str]) -> None:
    """Apply a ``YYYY-MM-DD`` display format to columns whose header
    contains the substring ``date`` (case-insensitive).

    Args:
        worksheet: The worksheet to format.
        header_names: Ordered list of header labels matching column order.
    """
    for col_index, header in enumerate(header_names, start=1):
        if any(hint in header.lower() for hint in _DATE_COLUMNS_HINT):
            column_letter = get_column_letter(col_index)
            for row in range(2, worksheet.max_row + 1):
                worksheet[f"{column_letter}{row}"].number_format = "YYYY-MM-DD"


def format_worksheet(worksheet: Worksheet, header_names: list[str]) -> None:
    """Apply the full standard formatting pipeline to a worksheet.

    Args:
        worksheet: The worksheet to format.
        header_names: Ordered list of header labels matching column order.
    """
    style_header_row(worksheet)
    freeze_header(worksheet)
    apply_autofilter(worksheet)
    apply_date_format(worksheet, header_names)
    autosize_columns(worksheet)
